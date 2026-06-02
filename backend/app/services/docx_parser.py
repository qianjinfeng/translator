"""Parse DOCX files into a list of DocumentSegment objects.

Uses python-docx to extract paragraphs, table cells, headers, and list items
while preserving formatting metadata (font, size, bold, italic, underline).
"""

from __future__ import annotations

from typing import Any

from docx import Document as DocxDocument
from docx.oxml.ns import qn
from docx.table import Table as DocxTable
from docx.text.paragraph import Paragraph
from docx.text.run import Run

from app.services.segment import DocumentSegment


def _extract_run_formatting(run: Run) -> dict[str, Any]:
    """Extract formatting metadata from a single run."""
    meta: dict[str, Any] = {}
    font = run.font
    if font.name:
        meta["font_name"] = font.name
    if font.size:
        meta["font_size"] = font.size.pt  # type: ignore[union-attr]
    if font.bold is not None:
        meta["bold"] = font.bold
    if font.italic is not None:
        meta["italic"] = font.italic
    if font.underline is not None:
        meta["underline"] = font.underline
    if font.color and font.color.rgb:
        meta["color"] = str(font.color.rgb)
    return meta


def _merge_formatting(runs_meta: list[dict[str, Any]]) -> dict[str, Any]:
    """Merge formatting metadata from multiple runs.  The first run's values are
    kept when a property is undefined in later runs (None)."""
    merged: dict[str, Any] = {}
    for m in runs_meta:
        for key, value in m.items():
            if key not in merged or value is not None:
                merged[key] = value
    return merged


def _paragraph_to_segment(
    para: Paragraph, index: int, segment_type: str = "paragraph"
) -> DocumentSegment:
    """Convert a python-docx Paragraph into a DocumentSegment."""
    runs_meta = [_extract_run_formatting(r) for r in para.runs if r.text.strip()]
    full_text = para.text

    # Determine heading level if any
    metadata: dict[str, Any] = {}
    if para.style and para.style.name and para.style.name.startswith("Heading"):
        parts = para.style.name.split()
        if len(parts) >= 2 and parts[1].isdigit():
            metadata["heading_level"] = int(parts[1])

    # List style detection
    if para.style and para.style.name:
        style_lower = para.style.name.lower()
        if "list" in style_lower:
            segment_type = "list_item"
            metadata["list_style"] = para.style.name

    metadata["segment_type"] = segment_type
    metadata.update(_merge_formatting(runs_meta))

    return DocumentSegment(
        index=index,
        source_text=full_text,
        segment_type=segment_type,
        metadata=metadata,
    )


def _extract_table_cells(table: DocxTable, start_index: int) -> list[DocumentSegment]:
    """Extract all cells from a table as individual segments.

    Iterates row-by-row, cell-by-cell in reading order.
    """
    segments: list[DocumentSegment] = []
    idx = start_index
    for row_idx, row in enumerate(table.rows):
        for col_idx, cell in enumerate(row.cells):
            text = cell.text.strip()
            if not text:
                continue
            cell_meta: dict[str, Any] = {
                "segment_type": "table_cell",
                "row": row_idx,
                "col": col_idx,
            }
            # Try to extract per-paragraph formatting within the cell
            for p in cell.paragraphs:
                if p.runs:
                    cell_meta.update(_extract_run_formatting(p.runs[0]))
                    break
            segments.append(
                DocumentSegment(
                    index=idx,
                    source_text=text,
                    segment_type="table_cell",
                    metadata=cell_meta,
                )
            )
            idx += 1
    return segments


def _is_list_paragraph(para: Paragraph) -> bool:
    """Check if a paragraph has list-like numbering or bullet numbering properties."""
    num_pr = para._element.find(qn("w:pPr"))
    if num_pr is not None:
        num_pr_elem = num_pr.find(qn("w:numPr"))
        if num_pr_elem is not None:
            return True
    return False


def parse_docx(file_path: str) -> list[DocumentSegment]:
    """Parse a DOCX file into a list of DocumentSegments.

    Args:
        file_path: Absolute path to the .docx file.

    Returns:
        A list of DocumentSegment objects in document order.
    """
    doc = DocxDocument(file_path)
    segments: list[DocumentSegment] = []
    # We walk the document body children in order, processing paragraphs and tables
    # as they appear in the XML body.
    body = doc.element.body

    index = 0
    for child in body:
        tag = child.tag.split("}")[-1] if "}" in child.tag else child.tag

        if tag == "p":
            # Find the paragraph object corresponding to this XML element
            para = _find_paragraph_by_element(doc.paragraphs, child)
            if para is None:
                continue
            text = para.text.strip()
            if not text:
                continue

            seg_type: str = "paragraph"
            if para.style and para.style.name and para.style.name.startswith("Heading"):
                seg_type = "header"
            elif _is_list_paragraph(para):
                seg_type = "list_item"

            seg = _paragraph_to_segment(para, index, seg_type)
            segments.append(seg)
            index += 1

        elif tag == "tbl":
            table = _find_table_by_element(doc.tables, child)
            if table is None:
                continue
            table_segments = _extract_table_cells(table, index)
            segments.extend(table_segments)
            index += len(table_segments)

    return segments


def _find_paragraph_by_element(
    paragraphs: list[Paragraph], element: object
) -> Paragraph | None:
    """Match a python-docx Paragraph to its underlying lxml element."""
    for p in paragraphs:
        if p._element is element:
            return p
    return None


def _find_table_by_element(tables: list[DocxTable], element: object) -> DocxTable | None:
    """Match a python-docx Table to its underlying lxml element."""
    for t in tables:
        if t._element is element:
            return t
    return None
