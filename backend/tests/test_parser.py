"""Tests for the unified DocumentParser interface."""

import os
import tempfile

import pytest

from app.services.parser import DocumentParser
from app.services.segment import DocumentSegment

# A reusable fixture helper


def _make_sample_docx(path: str) -> None:
    from docx import Document
    doc = Document()
    doc.add_paragraph("First paragraph")
    doc.add_paragraph("Second paragraph")
    doc.save(path)


class TestDocumentParser:
    def setup_method(self):
        self.parser = DocumentParser()

    def test_supported_formats(self):
        assert "docx" in self.parser.SUPPORTED_FORMATS
        assert "pdf" in self.parser.SUPPORTED_FORMATS

    def test_parse_docx(self):
        with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as f:
            path = f.name
        try:
            _make_sample_docx(path)
            segments = self.parser.parse(path, format="docx")
            assert len(segments) == 2
            assert all(isinstance(s, DocumentSegment) for s in segments)
        finally:
            os.unlink(path)

    def test_parse_unsupported_format(self):
        with pytest.raises(ValueError, match="Unsupported format"):
            self.parser.parse("/fake/path.txt", format="txt")

    def test_parse_format_case_insensitive(self):
        with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as f:
            path = f.name
        try:
            _make_sample_docx(path)
            segments_upper = self.parser.parse(path, format="DOCX")
            segments_lower = self.parser.parse(path, format="docx")
            assert len(segments_upper) == len(segments_lower)
        finally:
            os.unlink(path)

    def test_rebuild_docx_standalone(self):
        with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as f:
            orig_path = f.name
        out_path = orig_path.replace(".docx", "_out.docx")
        try:
            _make_sample_docx(orig_path)
            segments = self.parser.parse(orig_path, format="docx")
            trans_segments = [
                DocumentSegment(
                    index=s.index,
                    source_text=f"TR({s.source_text})",
                    segment_type=s.segment_type,
                    metadata=s.metadata,
                )
                for s in segments
            ]
            result_path = self.parser.rebuild(
                orig_path, trans_segments, out_path, mode="STANDALONE"
            )
            assert result_path == out_path
            assert os.path.exists(out_path)

            # Verify content
            from app.services.docx_parser import parse_docx
            result = parse_docx(out_path)
            assert result[0].source_text == "TR(First paragraph)"
        finally:
            for p in (orig_path, out_path):
                if os.path.exists(p):
                    os.unlink(p)

    def test_rebuild_docx_bilingual(self):
        with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as f:
            orig_path = f.name
        out_path = orig_path.replace(".docx", "_out.docx")
        try:
            _make_sample_docx(orig_path)
            segments = self.parser.parse(orig_path, format="docx")
            trans_segments = [
                DocumentSegment(
                    index=s.index,
                    source_text=f"TR({s.source_text})",
                    segment_type=s.segment_type,
                    metadata=s.metadata,
                )
                for s in segments
            ]
            _ = self.parser.rebuild(
                orig_path, trans_segments, out_path, mode="BILINGUAL"
            )
            assert os.path.exists(out_path)

            from app.services.docx_parser import parse_docx
            result = parse_docx(out_path)
            assert len(result) == 4
            assert result[0].source_text == "First paragraph"
            assert result[1].source_text == "TR(First paragraph)"
        finally:
            for p in (orig_path, out_path):
                if os.path.exists(p):
                    os.unlink(p)

    def test_rebuild_unsupported_format(self):
        from app.services.segment import DocumentSegment
        with pytest.raises(ValueError, match="Unsupported format"):
            self.parser.rebuild(
                "/fake/file.txt",
                [DocumentSegment(index=0, source_text="test")],
                "/fake/out.txt",
                mode="STANDALONE",
            )

    def test_rebuild_pdf_not_supported(self):
        with pytest.raises(NotImplementedError, match="Rebuilding PDFs"):
            self.parser.rebuild(
                "/fake/file.pdf",
                [DocumentSegment(index=0, source_text="test")],
                "/fake/out.pdf",
                mode="STANDALONE",
            )
