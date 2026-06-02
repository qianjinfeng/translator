"""Tests for the DOCX rebuilder."""

import os
import tempfile

from docx import Document

from app.services.docx_parser import parse_docx
from app.services.docx_rebuilder import rebuild_docx
from app.services.segment import DocumentSegment


def _make_simple_docx(path: str) -> None:
    """Create a minimal DOCX with two paragraphs."""
    doc = Document()
    doc.add_paragraph("Hello world")
    doc.add_paragraph("Second paragraph")
    doc.save(path)


def _make_table_docx(path: str) -> None:
    """Create a DOCX with a table between paragraphs."""
    doc = Document()
    doc.add_paragraph("Start")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Cell A1"
    table.cell(0, 1).text = "Cell B1"
    table.cell(1, 0).text = "Cell A2"
    table.cell(1, 1).text = "Cell B2"
    doc.add_paragraph("End")
    doc.save(path)


def _make_formatted_docx(path: str) -> None:
    """Create a DOCX with formatting to test preservation."""
    from docx.shared import Pt

    doc = Document()
    p = doc.add_paragraph()
    run = p.add_run("Formatted text")
    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(12)
    doc.save(path)


class TestDocxRebuilder:
    def test_standalone_replaces_text(self):
        with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as f:
            orig_path = f.name
        out_path = orig_path.replace(".docx", "_out.docx")
        try:
            _make_simple_docx(orig_path)
            segments = parse_docx(orig_path)
            # Modify segments with translated text
            trans_segments = [
                DocumentSegment(
                    index=s.index,
                    source_text=f"[TRANS]{s.source_text}[/TRANS]",
                    segment_type=s.segment_type,
                    metadata=s.metadata,
                )
                for s in segments
            ]
            rebuild_docx(orig_path, trans_segments, out_path, mode="STANDALONE")

            # Re-parse the output
            result = parse_docx(out_path)
            assert len(result) == 2
            assert result[0].source_text == "[TRANS]Hello world[/TRANS]"
            assert result[1].source_text == "[TRANS]Second paragraph[/TRANS]"
        finally:
            for p in (orig_path, out_path):
                if os.path.exists(p):
                    os.unlink(p)

    def test_bilingual_inserts_translation(self):
        with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as f:
            orig_path = f.name
        out_path = orig_path.replace(".docx", "_out.docx")
        try:
            _make_simple_docx(orig_path)
            segments = parse_docx(orig_path)
            trans_segments = [
                DocumentSegment(
                    index=s.index,
                    source_text=f"<TRANSLATED>{s.source_text}</TRANSLATED>",
                    segment_type=s.segment_type,
                    metadata=s.metadata,
                )
                for s in segments
            ]
            rebuild_docx(orig_path, trans_segments, out_path, mode="BILINGUAL")

            result = parse_docx(out_path)
            # Original paragraphs + translated paragraphs = 4
            assert len(result) == 4
            assert result[0].source_text == "Hello world"
            assert result[1].source_text == "<TRANSLATED>Hello world</TRANSLATED>"
            assert result[2].source_text == "Second paragraph"
            assert result[3].source_text == "<TRANSLATED>Second paragraph</TRANSLATED>"
        finally:
            for p in (orig_path, out_path):
                if os.path.exists(p):
                    os.unlink(p)

    def test_standalone_keeps_table_structure(self):
        with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as f:
            orig_path = f.name
        out_path = orig_path.replace(".docx", "_out.docx")
        try:
            _make_table_docx(orig_path)
            segments = parse_docx(orig_path)

            trans_segments = [
                DocumentSegment(
                    index=s.index,
                    source_text=f"T({s.source_text})",
                    segment_type=s.segment_type,
                    metadata=s.metadata,
                )
                for s in segments
            ]
            rebuild_docx(orig_path, trans_segments, out_path, mode="STANDALONE")

            result = parse_docx(out_path)
            # Start + 4 cells + End = 6
            assert len(result) == 6
            table_cells = [s for s in result if s.segment_type == "table_cell"]
            assert len(table_cells) == 4
            assert table_cells[0].source_text == "T(Cell A1)"
        finally:
            for p in (orig_path, out_path):
                if os.path.exists(p):
                    os.unlink(p)

    def test_bilingual_with_table(self):
        with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as f:
            orig_path = f.name
        out_path = orig_path.replace(".docx", "_out.docx")
        try:
            _make_table_docx(orig_path)
            segments = parse_docx(orig_path)

            trans_segments = [
                DocumentSegment(
                    index=s.index,
                    source_text=f"TR({s.source_text})",
                    segment_type=s.segment_type,
                    metadata=s.metadata,
                )
                for s in segments
            ]
            rebuild_docx(orig_path, trans_segments, out_path, mode="BILINGUAL")

            result = parse_docx(out_path)
            # Start + TR(Start) + 4 table cells (each has original + translation
            # appended inside the same cell) + End + TR(End) = 8
            assert len(result) == 8
            # First paragraph segments get bilingual inserts
            assert result[0].source_text == "Start"
            assert result[1].source_text == "TR(Start)"
            # Table cells have both original and translation (appended inside cell)
            table_cells = [s for s in result if s.segment_type == "table_cell"]
            assert len(table_cells) == 4
            assert "TR(Cell A1)" in table_cells[0].source_text
        finally:
            for p in (orig_path, out_path):
                if os.path.exists(p):
                    os.unlink(p)

    def test_output_path_returned(self):
        with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as f:
            orig_path = f.name
        out_path = orig_path.replace(".docx", "_out.docx")
        try:
            _make_simple_docx(orig_path)
            segments = parse_docx(orig_path)
            returned = rebuild_docx(orig_path, segments, out_path, mode="STANDALONE")
            assert returned == out_path
        finally:
            for p in (orig_path, out_path):
                if os.path.exists(p):
                    os.unlink(p)
